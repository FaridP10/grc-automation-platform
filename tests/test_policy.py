from datetime import date, timedelta
from pathlib import Path

import pytest
from docx import Document

from grc_platform.policy.catalog import load_policy_catalog
from grc_platform.policy.extract import extract_review_date, extract_text
from grc_platform.policy.gap_analysis import assess_all, assess_policy, coverage_summary

FIXTURES = Path(__file__).resolve().parents[1] / "data" / "fixtures"
MANIFEST_DIR = FIXTURES / "policies"


# --- catalog ---

def test_catalog_loads_12_policies():
    catalog = load_policy_catalog()
    assert len(catalog) == 12
    assert "access_control_policy" in catalog


# --- extract ---

def test_extract_text_md():
    text = extract_text(MANIFEST_DIR / "information_security_policy.md")
    assert "Management Commitment" in text


def test_extract_text_docx(tmp_path):
    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Last Reviewed: 2026-01-01")
    doc.add_paragraph("Acceptable use and prohibited use are both covered.")
    doc.save(str(path))
    text = extract_text(path)
    assert "Acceptable use" in text


def test_extract_text_unsupported_suffix(tmp_path):
    path = tmp_path / "policy.exe"
    path.write_text("nope")
    with pytest.raises(ValueError):
        extract_text(path)


def test_extract_review_date_iso_format():
    text = "Some heading\nLast Reviewed: 2026-03-01\nMore text"
    assert extract_review_date(text) == date(2026, 3, 1)


def test_extract_review_date_natural_language():
    text = "Effective Date: March 1, 2026"
    assert extract_review_date(text) == date(2026, 3, 1)


def test_extract_review_date_absent():
    assert extract_review_date("No date info here.") is None


# --- gap analysis ---

def test_assess_policy_missing_when_no_file():
    catalog = load_policy_catalog()
    gap = assess_policy(catalog["backup_policy"], None)
    assert gap.status == "missing"


def test_assess_policy_incomplete_missing_clause():
    catalog = load_policy_catalog()
    gap = assess_policy(catalog["access_control_policy"], MANIFEST_DIR / "access_control_policy.md")
    assert gap.status == "incomplete"
    assert "access review" in gap.detail


def test_assess_policy_stale_old_review_date():
    catalog = load_policy_catalog()
    gap = assess_policy(catalog["incident_response_policy"], MANIFEST_DIR / "incident_response_policy.md")
    assert gap.status == "stale"


def test_assess_policy_present():
    catalog = load_policy_catalog()
    gap = assess_policy(catalog["information_security_policy"], MANIFEST_DIR / "information_security_policy.md")
    assert gap.status == "present"


def test_assess_policy_stale_when_no_review_date_found(tmp_path):
    catalog = load_policy_catalog()
    path = tmp_path / "backup_policy.md"
    path.write_text("Backup frequency is daily. Retention is 30 days. Restore testing is quarterly.")
    gap = assess_policy(catalog["backup_policy"], path)
    assert gap.status == "stale"
    assert "No review" in gap.detail


def test_assess_all_and_coverage_summary():
    manifest = {
        "information_security_policy": MANIFEST_DIR / "information_security_policy.md",
        "access_control_policy": MANIFEST_DIR / "access_control_policy.md",
        "incident_response_policy": MANIFEST_DIR / "incident_response_policy.md",
        "acceptable_use_policy": MANIFEST_DIR / "acceptable_use_policy.docx",
    }
    gaps = assess_all(manifest)
    assert len(gaps) == 12

    summary = coverage_summary(gaps)
    assert summary["total"] == 12
    assert summary["missing"] == 8
    assert summary["incomplete"] == 1
    assert summary["stale"] == 1
    assert summary["present"] == 2
    assert summary["coverage_pct"] == round(100 * 2 / 12, 1)
